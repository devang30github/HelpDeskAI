"""Load documents (txt/md/pdf/docx), chunk, embed, and build the persistent index."""
import hashlib
import logging
import re
from datetime import datetime, timezone
from pathlib import Path

from app.core.config import Settings, get_settings
from app.core.constants import Category
from app.rag.chunking import chunk_document
from app.rag.embeddings import get_embedder
from app.rag.types import Document
from app.rag.vector_store import VectorStore, index_status, read_manifest

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}
VALID_CATEGORIES = {c.value for c in Category} | {"general"}
_HEADER = re.compile(r"^(Title|Category):\s*(.+?)\s*$", re.IGNORECASE)


def _slug(stem: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", stem.lower()).strip("-")


def _read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig", errors="replace")
    


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _read_docx(path: Path) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))  # paragraphs only; tables are not extracted
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


_READERS = {".txt": _read_txt, ".md": _read_txt, ".pdf": _read_pdf, ".docx": _read_docx}


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def parse_header(text: str) -> tuple[dict, str]:
    """Consume leading `Title:` / `Category:` lines; return (metadata, body)."""
    lines = text.split("\n")
    meta: dict = {}
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        m = _HEADER.match(line)
        if not m:
            break
        meta[m.group(1).lower()] = m.group(2)
        i += 1
    return meta, "\n".join(lines[i:]).strip()


def _candidate_files(kb_dir: Path) -> list[Path]:
    return sorted(
        p
        for p in Path(kb_dir).iterdir()
        if p.is_file() and not p.name.startswith((".", "_")) and not p.name.lower().startswith("readme")
    )
    
def list_kb_files(kb_dir: Path) -> list[Path]:
    """Supported knowledge-base files that the indexer will pick up."""
    kb_dir = Path(kb_dir)
    if not kb_dir.exists():
        return []
    return [p for p in _candidate_files(kb_dir) if p.suffix.lower() in SUPPORTED_EXTENSIONS]


def doc_id_for(filename: str) -> str:
    return _slug(Path(filename).stem)


def read_document_text(path: Path) -> str:
    ext = Path(path).suffix.lower()
    if ext not in _READERS:
        raise ValueError(f"unsupported extension '{ext}'")
    return _normalize(_READERS[ext](Path(path)))

def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fingerprint_kb(kb_dir: Path) -> dict[str, str]:
    """source_name -> sha256 for every supported file."""
    return {
        p.name: _sha256(p) for p in _candidate_files(kb_dir) if p.suffix.lower() in SUPPORTED_EXTENSIONS
    }


def load_documents(kb_dir: Path) -> tuple[list[Document], list[tuple[str, str]]]:
    """Returns (documents, skipped) where skipped is [(filename, reason)]."""
    kb_dir = Path(kb_dir)
    if not kb_dir.exists():
        raise FileNotFoundError(f"Knowledge base directory not found: {kb_dir}")

    docs: list[Document] = []
    skipped: list[tuple[str, str]] = []
    seen: dict[str, str] = {}

    for path in _candidate_files(kb_dir):
        ext = path.suffix.lower()
        if ext not in SUPPORTED_EXTENSIONS:
            skipped.append((path.name, f"unsupported extension '{ext}'"))
            continue
        try:
            raw = _normalize(_READERS[ext](path))
        except Exception as e:  # corrupt/encrypted files must not abort the whole build
            skipped.append((path.name, f"read error: {e}"))
            continue

        meta, body = parse_header(raw)
        if not body:
            skipped.append((path.name, "no extractable text"))
            continue

        doc_id = _slug(path.stem)
        if not doc_id:
            skipped.append((path.name, "filename produces an empty doc_id"))
            continue
        if doc_id in seen:
            raise ValueError(f"Duplicate doc_id '{doc_id}' from '{seen[doc_id]}' and '{path.name}'")
        seen[doc_id] = path.name

        category = meta.get("category")
        if category and category not in VALID_CATEGORIES:
            logger.warning("Unknown category '%s' in %s; ignoring it", category, path.name)
            meta.pop("category")

        docs.append(
            Document(
                doc_id=doc_id,
                source_name=path.name,
                text=body,
                sha256=_sha256(path),
                metadata={"title": meta.get("title", path.stem), **({"category": meta["category"]} if "category" in meta else {})},
            )
        )
    return docs, skipped


def build_index(
    settings: Settings | None = None,
    chunk_size: int | None = None,
    chunk_overlap: int | None = None,
) -> dict:
    s = settings or get_settings()
    size = chunk_size if chunk_size is not None else s.chunk_size
    overlap = chunk_overlap if chunk_overlap is not None else s.chunk_overlap

    docs, skipped = load_documents(s.kb_dir)
    if not docs:
        raise RuntimeError(f"No usable documents found in {s.kb_dir}")

    chunks = []
    for doc in docs:
        chunks.extend(chunk_document(doc, size, overlap))
    logger.info("Loaded %d documents -> %d chunks", len(docs), len(chunks))

    embedder = get_embedder(s.embedding_model)
    vectors = embedder.encode([c.text for c in chunks])

    per_doc: dict[str, dict] = {}
    for c in chunks:
        entry = per_doc.setdefault(
            c.doc_id,
            {"source_name": c.source_name, "title": c.metadata.get("title"),
             "category": c.metadata.get("category"), "num_chunks": 0},
        )
        entry["num_chunks"] += 1

    manifest = {
        "embedding_model": s.embedding_model,
        "chunk_size": size,
        "chunk_overlap": overlap,
        "num_documents": len(docs),
        "built_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "files": {d.source_name: d.sha256 for d in docs},
        "documents": per_doc,
        "skipped": [{"file": f, "reason": r} for f, r in skipped],
    }
    store = VectorStore.build(chunks, vectors, manifest)
    store.save(s.index_dir)

    # A running process may hold an old retriever; clear it.
    from app.rag.retrieval import reset_retriever_cache

    reset_retriever_cache()
    return store.manifest


def check_index(settings: Settings | None = None) -> tuple[bool, list[str]]:
    """Returns (is_stale, reasons). Detects added/removed/changed files and config drift."""
    s = settings or get_settings()
    if not index_status(s.index_dir)["built"]:
        return True, ["index has not been built"]

    manifest = read_manifest(s.index_dir)
    reasons: list[str] = []
    old, new = manifest.get("files", {}), fingerprint_kb(s.kb_dir)
    for name in sorted(new.keys() - old.keys()):
        reasons.append(f"new file: {name}")
    for name in sorted(old.keys() - new.keys()):
        reasons.append(f"removed file: {name}")
    for name in sorted(new.keys() & old.keys()):
        if new[name] != old[name]:
            reasons.append(f"changed file: {name}")
    for key, current in (
        ("embedding_model", s.embedding_model),
        ("chunk_size", s.chunk_size),
        ("chunk_overlap", s.chunk_overlap),
    ):
        if manifest.get(key) != current:
            reasons.append(f"{key} changed: index={manifest.get(key)!r}, config={current!r}")
    return bool(reasons), reasons