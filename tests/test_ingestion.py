import pytest
from docx import Document as DocxDocument

from app.rag.ingestion import load_documents, parse_header


def test_parse_header():
    meta, body = parse_header("Title: A Guide\nCategory: network_wifi\n\nBody text here.")
    assert meta == {"title": "A Guide", "category": "network_wifi"}
    assert body == "Body text here."


def test_parse_header_absent():
    meta, body = parse_header("Just text.\nMore text.")
    assert meta == {} and body == "Just text.\nMore text."


def test_loads_txt_and_docx_and_skips_unsupported(tmp_path):
    (tmp_path / "WiFi Guide.txt").write_text(
        "Title: Wi-Fi\nCategory: network_wifi\n\nReset the router.", encoding="utf-8")
    d = DocxDocument()
    d.add_paragraph("Title: Printing")
    d.add_paragraph("Category: printing_office_equipment")
    d.add_paragraph("Clear the paper jam.")
    d.save(str(tmp_path / "printing.docx"))
    (tmp_path / "notes.xyz").write_text("ignored")
    (tmp_path / "README.md").write_text("not a knowledge document")

    docs, skipped = load_documents(tmp_path)
    by_id = {doc.doc_id: doc for doc in docs}
    assert set(by_id) == {"wifi-guide", "printing"}
    assert by_id["printing"].metadata["category"] == "printing_office_equipment"
    assert by_id["printing"].text == "Clear the paper jam."
    assert [name for name, _ in skipped] == ["notes.xyz"]


def test_unknown_category_is_dropped(tmp_path):
    (tmp_path / "a.txt").write_text("Title: A\nCategory: made_up\n\nSome body text.")
    docs, _ = load_documents(tmp_path)
    assert "category" not in docs[0].metadata


def test_duplicate_doc_ids_are_rejected(tmp_path):
    (tmp_path / "a-b.txt").write_text("Some text one.")
    (tmp_path / "a_b.txt").write_text("Some text two.")
    with pytest.raises(ValueError):
        load_documents(tmp_path)


def test_loads_pdf(tmp_path):
    fpdf = pytest.importorskip("fpdf")
    pdf = fpdf.FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(0, 6, "Title: Facilities\nCategory: facilities_maintenance\n\nReport leaks immediately.")
    pdf.output(str(tmp_path / "fac.pdf"))
    docs, _ = load_documents(tmp_path)
    assert docs[0].doc_id == "fac"
    assert "Report leaks immediately." in docs[0].text